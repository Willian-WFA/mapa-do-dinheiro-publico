import json,csv,collections,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
X='/media/hitsuzen/Novo volume/money/arrecadacao_despesas_2019_2026.xlsx'
res=json.load(open('cvm_pos.json'))
# estoque interno de mercado (Tesouro)
est=collections.defaultdict(float)
with open('/media/hitsuzen/Novo volume/money/dados_portal/tesouro/estoque-da-divida-publica-federal/estoquedpf.csv',encoding='latin1',newline='') as f:
    for r in csv.DictReader(f,delimiter=';'):
        if r['Classe da Carteira']=='Mercado' and 'Interna' in r['Tipo de Divida']:
            try: est[r['Mes do Estoque']]+=float(r['Valor do Estoque'].replace('.','').replace(',','.'))
            except ValueError: pass
def key(p): y,m=p.split('-'); return f'{m}/{y}'
wb=openpyxl.load_workbook(X)
if 'Carteiras de fundos' in wb.sheetnames: del wb['Carteiras de fundos']
ws=wb.create_sheet('Carteiras de fundos',3)
hdr=PatternFill('solid',fgColor='1F3864'); grp=PatternFill('solid',fgColor='D9E1F2')
ws['A1']='Quem detém títulos públicos federais por meio de fundos de investimento (carteiras próprias, CVM)'; ws['A1'].font=Font(bold=True,size=14)
ws['A2']=('Fonte: CVM, Composição e Diversificação das Aplicações (CDA), bloco de títulos públicos, posição final em dezembro (2026 = agosto), agrupada pela gestora do fundo (cadastro CVM). '
          'Não inclui compromissadas, carteira própria de bancos, previdência fechada fora de fundos, não residentes nem pessoas físicas diretas. É um recorte parcial: valores em R$ bilhões.')
ws['A2'].alignment=Alignment(wrap_text=True); ws['A2'].font=Font(italic=True,size=9); ws.row_dimensions[2].height=44
per=sorted(res)
for c,v in enumerate(['Item']+[p[:4]+(' (ago)' if p=='2026-08' else '') for p in per],1):
    x=ws.cell(row=4,column=c,value=v); x.fill=hdr; x.font=Font(bold=True,color='FFFFFF'); x.alignment=Alignment(horizontal='center')
r=5
def line(label,vals,fmt='#,##0',bold=False):
    global r
    ws.cell(row=r,column=1,value=label).font=Font(bold=bold)
    for j,v in enumerate(vals):
        if v is not None: ws.cell(row=r,column=2+j,value=v).number_format=fmt
    r+=1
line('Títulos públicos nas carteiras de fundos (R$ bi)',[res[p]['total']/1e9 for p in per],bold=True)
line('Operações compromissadas (R$ bi, não somadas)',[res[p]['compromissadas']/1e9 for p in per])
line('Estoque da dívida interna em mercado (R$ bi, Tesouro)',[ (est.get(key(p))/1e9 if est.get(key(p)) else None) for p in per])
line('Fundos / dívida interna em mercado (%)',[ (res[p]['total']/est[key(p)]*100 if est.get(key(p)) else None) for p in per],'0.0')
line('Nº de grupos gestores com posição',[res[p]['n_gestores'] for p in per])
r+=1
ws.cell(row=r,column=1,value='POR GESTORA (grupo econômico) – R$ bi').font=Font(bold=True)
for c in range(1,len(per)+2): ws.cell(row=r,column=c).fill=grp
r+=1
names=[]
for p in per:
    for g,_ in res[p]['top'][:10]:
        if g not in names: names.append(g)
for g in names:
    line(g,[ next((v/1e9 for n,v in res[p]['top'] if n==g),None) for p in per],'#,##0.0')
r+=1
ws.cell(row=r,column=1,value='Participação das 5 maiores gestoras no total dos fundos (%)').font=Font(bold=True)
for j,p in enumerate(per):
    top5=sum(v for g,v in res[p]['top'] if g!='Sem gestor')
    top5=sum(v for g,v in sorted([(g,v) for g,v in res[p]['top'] if g!='Sem gestor'],key=lambda x:-x[1])[:5])
    ws.cell(row=r,column=2+j,value=top5/res[p]['total']*100).number_format='0.0'
r+=2
ws.cell(row=r,column=1,value='Tipos de título (R$ bi, último período)').font=Font(bold=True); r+=1
for t,v in res[per[-1]]['tipo']:
    ws.cell(row=r,column=1,value=t.title()); ws.cell(row=r,column=2,value=v/1e9).number_format='#,##0.0'; r+=1
ws.column_dimensions['A'].width=58
for j in range(len(per)): ws.column_dimensions[openpyxl.utils.get_column_letter(2+j)].width=11
ws.freeze_panes='B5'
wb.save(X); print('ok')
print([ (p, round(res[p]['total']/1e9), round(est.get(key(p),0)/1e9)) for p in per])
