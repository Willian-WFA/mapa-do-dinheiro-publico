import json,os,csv,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
ROOT='/media/hitsuzen/Novo volume/money'; DP=ROOT+'/dados_portal'
X=ROOT+'/arrecadacao_despesas_2019_2026.xlsx'
dg=json.load(open('dgt_series.json'))
wb=openpyxl.load_workbook(X)
for n in ('Gastos tributários (DGT)','Fontes e arquivos'):
    if n in wb.sheetnames: del wb[n]
hdr=PatternFill('solid',fgColor='1F3864'); grp=PatternFill('solid',fgColor='D9E1F2')
# ---------- DGT ----------
ws=wb.create_sheet('Gastos tributários (DGT)',wb.sheetnames.index('Top50 isenção') if 'Top50 isenção' in wb.sheetnames else None)
ws['A1']='Gastos tributários oficiais – Demonstrativo da Receita Federal (DGT), R$ bilhões nominais'; ws['A1'].font=Font(bold=True,size=14)
ws['A2']='Efetivo = estimativa das bases efetivas (edição mais recente em que o ano é efetivo). Projeção = edições do DGT. PLOA = estimativa feita na época do projeto de lei orçamentária. Arquivos baixados em dados_portal/dgt/; links na aba "Fontes e arquivos".'
ws['A2'].font=Font(italic=True,size=9)
Y=list(range(2013,2027))
for c,v in enumerate(['Item']+[str(y) for y in Y],1):
    x=ws.cell(row=4,column=c,value=v); x.fill=hdr; x.font=Font(bold=True,color='FFFFFF'); x.alignment=Alignment(horizontal='center')
r=5
def line(lab,d,bold=False):
    global r
    ws.cell(row=r,column=1,value=lab).font=Font(bold=bold)
    for j,y in enumerate(Y):
        v=d.get(str(y))
        if v is not None: ws.cell(row=r,column=2+j,value=round(v,1)).number_format='#,##0.0'
    r+=1
line('TOTAL – bases efetivas (oficial)',dg['total'],True)
line('   projeção das edições do DGT',dg['proj'])
line('   previsão da PLOA do ano',dg['ploa_total'])
pib=dg['pib']
line('   % do PIB (efetivo)',{y:dg['total'][y]*1e9/pib[y]*100 for y in dg['total'] if y in pib} if all(k in pib for k in ['2021']) else {})
r+=1
ws.cell(row=r,column=1,value='POR FUNÇÃO ORÇAMENTÁRIA (efetivo, R$ bi)').font=Font(bold=True)
for c in range(1,len(Y)+2): ws.cell(row=r,column=c).fill=grp
r+=1
funs=sorted(dg['fun'].items(),key=lambda kv:-max(kv[1].values()) if kv[1] else 0)
for n,d in funs: line(n,d)
r+=1
ws.cell(row=r,column=1,value='Observação: a série oficial usa metodologia revisada a cada edição; por isso o valor efetivo de um ano pode ser maior que a previsão da PLOA daquele ano (ex.: 2021 = R$ 421 bi efetivo vs R$ 308 bi na PLOA).').font=Font(italic=True,size=9)
ws.column_dimensions['A'].width=44
for j in range(len(Y)): ws.column_dimensions[openpyxl.utils.get_column_letter(2+j)].width=9
ws.freeze_panes='B5'
# ---------- Fontes e arquivos ----------
def folder_info(p):
    n=0;s=0
    for dp,_,fs in os.walk(p):
        for f in fs:
            n+=1; s+=os.path.getsize(os.path.join(dp,f))
    return n,s
rows=[]
def add(tema,desc,local,link,extra=''):
    n,s=(0,0)
    if local and os.path.exists(ROOT+'/'+local):
        pth=ROOT+'/'+local
        n,s=folder_info(pth) if os.path.isdir(pth) else (1,os.path.getsize(pth))
    rows.append([tema,desc,local,n,round(s/1e6,1),link,extra])
cat=json.load(open(ROOT+'/catalogo_portal.json'))
titles={'servidores':'Servidores civis e militares do Executivo (remuneração, cadastro)','renuncias':'Renúncias fiscais por empresa','emendas-parlamentares':'Emendas parlamentares','emendas-parlamentares-documentos':'Emendas parlamentares por documento','despesas-execucao':'Execução da despesa por órgão/função/ação','orcamento-despesa':'Orçamento da despesa (anual)','receitas':'Receitas públicas','bpc':'BPC – benefícios pagos','bolsa-familia-pagamentos':'Bolsa Família – pagamentos','novo-bolsa-familia':'Novo Bolsa Família','auxilio-brasil':'Auxílio Brasil','auxilio-emergencial':'Auxílio Emergencial','pe-de-meia':'Pé-de-Meia','gas-do-povo':'Gás do Povo','viagens':'Viagens a serviço','imoveis-funcionais':'Imóveis funcionais','cpcc':'Cartões de compras centralizadas','ceis':'CEIS – empresas inidôneas','cnep':'CNEP – empresas punidas','ceaf':'CEAF – servidores expulsos','cepim':'CEPIM – entidades impedidas','convenios':'Convênios','acordos-leniencia':'Acordos de leniência','pep':'Pessoas expostas politicamente','apoiamento-emendas-parlamentares-documentos':'Apoiamento de emendas','auxilio-reconstrucao':'Auxílio Reconstrução'}
for d in sorted(os.listdir(DP)):
    p=DP+'/'+d
    if os.path.isdir(p) and d in cat:
        add('Portal da Transparência',titles.get(d,d),'dados_portal/'+d,f'https://portaldatransparencia.gov.br/download-de-dados/{d}','download por mês/ano: /download-de-dados/'+d+'/AAAAMM')
# DGT
for r_ in csv.DictReader(open(DP+'/dgt/INDICE_DGT.csv',encoding='utf8')):
    add('Receita Federal – DGT',r_['documento'],('dados_portal/dgt/'+r_['arquivo_local']) if r_['arquivo_local'] else '',r_['link_download'],'página: '+r_['link_pagina'])
add('Tesouro Nacional','Resultado do Tesouro Nacional – série histórica (1997–2026)','dados_portal/outros/RTN_serie_historica_jul2026.xlsx','https://www.tesourotransparente.gov.br/ckan/dataset/resultado-do-tesouro-nacional','arquivo usado: seriehistoricajul26.xlsx')
for ds in sorted(os.listdir(DP+'/tesouro')):
    add('Tesouro Nacional (dados abertos)',ds.replace('-',' '),'dados_portal/tesouro/'+ds,'https://www.tesourotransparente.gov.br/ckan/dataset/'+ds)
add('Tesouro Nacional','Relatório Mensal da Dívida (PDF; detentores por categoria)','dados_portal/rmd','https://www.tesourotransparente.gov.br/publicacoes/relatorio-mensal-da-divida-rmd')
add('CVM','Carteiras de fundos (CDA) – títulos públicos','dados_portal/cvm','https://dados.cvm.gov.br/dataset/fi-doc-cda','arquivos hist_AAAA.zip e m_AAAAMM.zip')
add('CVM','Cadastro de fundos e classes','dados_portal/cvm/cad_fi.csv','https://dados.cvm.gov.br/dados/FI/CAD/DADOS/cad_fi.csv','registro_fundo_classe.zip no mesmo dataset')
add('DadosJusBr','Remuneração do Judiciário (pacotes por tribunal e mês)','dados_portal/judiciario','https://dadosjusbr.org','API: https://api.dadosjusbr.org/v2/orgaos')
add('Câmara dos Deputados','Execução orçamentária (2015–2025)','dados_portal/outros/camara','https://www2.camara.leg.br/transparencia/receitas-e-despesas/despesas/despesas-consolidadas')
add('Senado Federal','Dotação autorizada e despesas executadas','dados_portal/outros/senado','https://www12.senado.leg.br/dados-abertos/administrativo/orcamento-do-senado/dotacao-e-despesas/dotacao-autorizada-e-despesas-executadas')
add('Banco Central','Séries SGS (juros nominais, resultado primário, PIB, IPCA)','dados_portal/outros/bcb','https://api.bcb.gov.br/dados/serie/bcdata.sgs.{série}/dados?formato=json','séries 5727, 5793, 4382, 433')
add('Receita Federal','Carga tributária no Brasil (tabelas 2024)','dados_portal/outros/receita_ctb','https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/estudos/carga-tributaria')
add('IBGE','População residente estimada (SIDRA tabela 6579)','','https://apisidra.ibge.gov.br/values/t/6579/n1/all/v/9324/p/all','consultado via API; valores em dados_portal/outros/bcb/hist_ext.json')
add('Banco Mundial / UNESCO','Matrículas por nível (SE.PRM/SEC/TER/PRE.ENRL)','dados_portal/outros/bcb/wb_enrl.json','https://api.worldbank.org/v2/country/BRA/indicator/SE.PRM.ENRL?format=json')
add('INEP','Sinopses estatísticas da educação básica e superior','', 'https://www.gov.br/inep/pt-br/acesso-a-informacao/dados-abertos/sinopses-estatisticas','arquivos muito grandes (>100 MB cada); baixados e apagados após extrair totais; servidor com certificado não confiável')
add('ANS','Beneficiários de planos (informações consolidadas)','', 'https://dadosabertos.ans.gov.br/FTP/PDA/informacoes_consolidadas_de_beneficiarios-024/','não baixado; valores aproximados na planilha')
ws2=wb.create_sheet('Fontes e arquivos')
ws2['A1']='Onde estão os arquivos baixados e de onde vieram'; ws2['A1'].font=Font(bold=True,size=14)
ws2['A2']='Pasta base: '+DP+' . Cada linha: arquivo(s) local(is) e o link para baixar da fonte.'; ws2['A2'].font=Font(italic=True,size=9)
for c,v in enumerate(['Fonte','Conteúdo','Local (relativo à pasta do projeto)','Nº de arquivos','Tamanho (MB)','Link da fonte','Observação'],1):
    x=ws2.cell(row=4,column=c,value=v); x.fill=hdr; x.font=Font(bold=True,color='FFFFFF')
for i,rw in enumerate(rows,5):
    for c,v in enumerate(rw,1):
        x=ws2.cell(row=i,column=c,value=v)
        if c==6 and str(v).startswith('http'): x.hyperlink=v; x.font=Font(color='0563C1',underline='single')
for c,w in zip('ABCDEFG',[26,52,52,10,12,70,50]): ws2.column_dimensions[c].width=w
ws2.freeze_panes='A5'; ws2.auto_filter.ref=f'A4:G{4+len(rows)}'
wb.save(X)
csv.writer(open(ROOT+'/FONTES_E_ARQUIVOS.csv','w',newline='',encoding='utf-8-sig'),delimiter=';').writerows([['Fonte','Conteúdo','Local','Nº arquivos','MB','Link da fonte','Observação']]+rows)
print('ok',len(rows),'fontes')
