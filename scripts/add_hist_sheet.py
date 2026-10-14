import json,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
X='/media/hitsuzen/Novo volume/money/arrecadacao_despesas_2019_2026.xlsx'
H=json.load(open('hist_rows.json')); Y=H['years']; fac={int(k):v for k,v in H['fac'].items()}
wb=openpyxl.load_workbook(X)
for n in ('Série histórica','Série histórica (real)'):
    if n in wb.sheetnames: del wb[n]
hdr=PatternFill('solid',fgColor='1F3864'); grp=PatternFill('solid',fgColor='D9E1F2')
def mk(name,real,pos):
    ws=wb.create_sheet(name,pos)
    ws['A1']='Série histórica 1997–2026 – '+('valores REAIS (R$ de 2025, IPCA)' if real else 'valores NOMINAIS')+' (R$ bilhões, salvo indicação)'
    ws['A1'].font=Font(bold=True,size=14)
    ws['A2']='Cada série começa no primeiro ano com registro na fonte. 2026 parcial. Fontes: Tesouro (RTN), Portal da Transparência, Banco Central, CVM, IBGE, Banco Mundial, Câmara e Senado.'
    ws['A2'].font=Font(italic=True,size=9)
    ws.cell(row=4,column=1,value='Item').fill=hdr; ws.cell(row=4,column=1).font=Font(bold=True,color='FFFFFF')
    for j,y in enumerate(Y):
        c=ws.cell(row=3,column=2+j,value=H['gov'][str(y)].replace(' (parcial)','')); c.font=Font(size=8,italic=True); c.alignment=Alignment(horizontal='center')
        c=ws.cell(row=4,column=2+j,value=str(y)+('*' if y==2026 else '')); c.fill=hdr; c.font=Font(bold=True,color='FFFFFF'); c.alignment=Alignment(horizontal='center')
    r=5; cur=None
    for row in H['rows']:
        if row['sec']!=cur:
            cur=row['sec']; ws.cell(row=r,column=1,value=cur).font=Font(bold=True)
            for c in range(1,len(Y)+2): ws.cell(row=r,column=c).fill=grp
            r+=1
        ws.cell(row=r,column=1,value=row['label'])
        money=row['kind'] in ('val','stock','pp')
        for j,y in enumerate(Y):
            v=row['data'].get(str(y))
            if v is None: continue
            if real and money: v=v*fac[y]
            c=ws.cell(row=r,column=2+j,value=round(v,2)); c.number_format='#,##0.0' if abs(v)<1000 else '#,##0'
        r+=1
    ws.column_dimensions['A'].width=72
    for j in range(len(Y)): ws.column_dimensions[openpyxl.utils.get_column_letter(2+j)].width=9.5
    ws.freeze_panes='B5'
mk('Série histórica (real)',True,1); mk('Série histórica',False,2)
wb.save(X); print('ok')
