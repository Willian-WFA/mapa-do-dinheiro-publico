import json,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
X='/media/hitsuzen/Novo volume/money/arrecadacao_despesas_2019_2026.xlsx'
wb=openpyxl.load_workbook(X)
hdr=PatternFill('solid',fgColor='1F3864'); grp=PatternFill('solid',fgColor='D9E1F2')
if 'Complementos 2016-2019' in wb.sheetnames: del wb['Complementos 2016-2019']
ws=wb.create_sheet('Complementos 2016-2019')
ws['A1']='Complementos: carga tributária por esfera, juros do setor público e militares (despesa paga)'; ws['A1'].font=Font(bold=True,size=13)
def head(r,vals):
    for c,v in enumerate(vals,1):
        x=ws.cell(row=r,column=c,value=v); x.font=Font(bold=True,color='FFFFFF'); x.fill=hdr; x.alignment=Alignment(horizontal='center')
yrs=list(range(2016,2025))
# 1. esfera
sh={'União':[.662,.6578,.6561,.6491,.6404,.643,.6582,.6583,.6614],'Estados':[.2718,.2748,.274,.2767,.2879,.2851,.2707,.2651,.2628],'Municípios':[.0661,.0674,.0699,.0741,.0717,.0719,.0712,.0766,.0759]}
tot={2017:1977.7,2018:2136.3,2019:2249.4,2020:2206.1,2021:2771.4,2022:3144.1,2023:3307.6,2024:3782.1}
ws['A3']='1) Arrecadação tributária bruta por esfera (R$ bi) – Receita Federal, Carga Tributária no Brasil 2024 (participação × total)'; ws['A3'].font=Font(bold=True)
head(4,['Esfera']+[str(y) for y in yrs])
r=5
for k,v in sh.items():
    ws.cell(row=r,column=1,value=k)
    for j,y in enumerate(yrs):
        if y in tot:
            c=ws.cell(row=r,column=2+j,value=round(v[j]*tot[y],1)); c.number_format='#,##0.0'
    r+=1
ws.cell(row=r,column=1,value='Total').font=Font(bold=True)
for j,y in enumerate(yrs):
    if y==2016: ws.cell(row=r,column=2+j,value=2021.0).number_format='#,##0.0'
    elif y in tot: ws.cell(row=r,column=2+j,value=tot[y]).number_format='#,##0.0'
r+=1
ws.cell(row=r,column=1,value='2016: total (R$ 2,021 tri) e por esfera (estados 514,5; União 1.381; municípios 124,8) são da edição de 2016 do estudo, com metodologia anterior; não comparar diretamente com 2017+.').font=Font(italic=True,size=9)
ws.cell(row=r+1,column=1,value='Estados 2016 (edição 2016): 514,5 | União 1.381 | Municípios 124,8').font=Font(italic=True,size=9)
r+=3
# 2. juros BC
J={2016:407.5,2017:401.1,2018:378.9,2019:367.2,2020:312.0,2021:448.8,2022:586.6,2023:717.9,2024:950.6,2025:1007.6}
ws.cell(row=r,column=1,value='2) Juros nominais do setor público consolidado (R$ bi) – Banco Central, séries SGS 5727, 5793 e 4382 (dezembro, 12 meses)').font=Font(bold=True); r+=1
yj=list(range(2016,2026)); head(r,['Item']+[str(y) for y in yj]); r+=1
ws.cell(row=r,column=1,value='Juros nominais')
for j,y in enumerate(yj): ws.cell(row=r,column=2+j,value=J[y]).number_format='#,##0.0'
r+=3
# 3. militares
try: D=json.load(open('desp_cache.json'))
except Exception: D={}
ws.cell(row=r,column=1,value='3) Militares federais – despesa paga no orçamento (R$ bi), ações "Ativos/Inativos/Pensões Militares da União" – Portal da Transparência, Despesas').font=Font(bold=True); r+=1
head(r,['Grupo']+[str(y) for y in range(2016,2020)]); r+=1
byy={}
for k,v in D.items():
    y=int(k[:4])
    for g,x in v.items(): byy.setdefault(y,{}).setdefault(g,[0,0]); byy[y][g][0]+=x; byy[y][g][1]+=1
for g in ['Ativa','Inativos (reserva/reforma)','Pensionistas']:
    ws.cell(row=r,column=1,value=g)
    for j,y in enumerate(range(2016,2020)):
        if y in byy and g in byy[y]:
            c=ws.cell(row=r,column=2+j,value=round(byy[y][g][0]/1e9,2)); c.number_format='#,##0.00'
    r+=1
ws.cell(row=r,column=1,value='Meses carregados')
for j,y in enumerate(range(2016,2020)):
    if y in byy: ws.cell(row=r,column=2+j,value=max(v[1] for v in byy[y].values()))
r+=1
ws.cell(row=r,column=1,value='Quantidade de pessoas em reserva e pensionistas antes de 2020: não publicada pelo Portal (o TCU determinou a divulgação em set/2019).').font=Font(italic=True,size=9)
ws.column_dimensions['A'].width=30
for c in range(2,13): ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width=12
m=wb['Militares']
for row in m.iter_rows(min_row=5):
    y=row[0].value
    if y in byy:
        for col,g in ((7,'Inativos (reserva/reforma)'),(10,'Pensionistas')):
            if m.cell(row=row[0].row,column=col).value is None and g in byy[y]:
                c=m.cell(row=row[0].row,column=col,value=round(byy[y][g][0]/1e9,2)); c.number_format='#,##0.00'
        m.cell(row=row[0].row,column=14,value='reserva/pensões: despesa paga do orçamento (qtde n/d)')
wb.save(X); print('ok',{y:{g:round(v[0]/1e9,2) for g,v in d.items()} for y,d in byy.items()})
