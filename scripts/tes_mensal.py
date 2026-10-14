import openpyxl,collections,json
wb=openpyxl.load_workbook('/media/hitsuzen/Novo volume/money/dados_portal/tesouro/despesas-mensais-2008/despesas_mensais_2008.xlsx',read_only=True,data_only=True)
ws=wb[wb.sheetnames[0]]
it=ws.iter_rows(values_only=True); h=next(it)
iF=h.index('Funcao_nome'); iA=h.index('Ano'); iV=h.index('Valor'); iM=h.index('Mes')
out=collections.defaultdict(lambda:collections.defaultdict(float)); meses=collections.defaultdict(set); n=0
for r in it:
    try: out[int(r[iA])][r[iF]]+=float(r[iV] or 0); meses[int(r[iA])].add(int(r[iM]))
    except (TypeError,ValueError): pass
    n+=1
json.dump({'func':{y:dict(d) for y,d in out.items()},'meses':{y:sorted(m) for y,m in meses.items()}},open('tes_mensal_func.json','w'),ensure_ascii=False)
print(n,{y:(round(sum(d.values())/1e9),len(meses[y])) for y,d in sorted(out.items())})
for y in (2008,2012,2018): print(y,{k[:22]:round(v/1e9,1) for k,v in sorted(out[y].items(),key=lambda x:-x[1])[:8]})
