import json,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
X='/media/hitsuzen/Novo volume/money/arrecadacao_despesas_2019_2026.xlsx'
T=json.load(open('top50.json')); Y=list(range(2016,2027))
wb=openpyxl.load_workbook(X)
if 'Top50 ano a ano' in wb.sheetnames: del wb['Top50 ano a ano']
ws=wb.create_sheet('Top50 ano a ano',wb.sheetnames.index('Top50 isenção'))
hdr=PatternFill('solid',fgColor='1F3864'); na=PatternFill('solid',fgColor='EDEDED'); alt=PatternFill('solid',fgColor='F2F2F2')
ws['A1']='50 empresas com maior isenção fiscal – ano a ano (valor em R$ milhões)'; ws['A1'].font=Font(bold=True,size=14)
ws['A2']='Fonte: Portal da Transparência – Renúncias Fiscais (valores declarados). 2024 é parcial; 2025 e 2026 ainda não foram publicados pelo portal (colunas vazias). Matriz e filiais somadas pela raiz do CNPJ.'
ws['A2'].font=Font(italic=True,size=9)
ws.cell(row=4,column=1,value='#').font=Font(bold=True)
for j,y in enumerate(Y):
    c0=2+j*2
    lab=str(y)+(' (parcial)' if y==2024 else ' (não publicado)' if y>2024 else '')
    ws.cell(row=3,column=c0,value=lab)
    ws.merge_cells(start_row=3,start_column=c0,end_row=3,end_column=c0+1)
    for c in (c0,c0+1):
        ws.cell(row=3,column=c).fill=hdr if y<=2024 else PatternFill('solid',fgColor='7F7F7F')
        ws.cell(row=3,column=c).font=Font(bold=True,color='FFFFFF'); ws.cell(row=3,column=c).alignment=Alignment(horizontal='center')
    ws.cell(row=4,column=c0,value='Empresa').font=Font(bold=True)
    ws.cell(row=4,column=c0+1,value='R$ mi').font=Font(bold=True)
    t=T.get(str(y))
    for i in range(50):
        r=5+i
        if i==0 or True: ws.cell(row=r,column=1,value=i+1)
        if t and i<len(t['top']):
            e=t['top'][i]
            ws.cell(row=r,column=c0,value=e['nome'].title()[:40])
            x=ws.cell(row=r,column=c0+1,value=round(e['valor']/1e6,1)); x.number_format='#,##0.0'
        elif not t:
            ws.cell(row=r,column=c0).fill=na; ws.cell(row=r,column=c0+1).fill=na
    # totais
    if t:
        ws.cell(row=56,column=c0,value='Soma das 50').font=Font(bold=True)
        ws.cell(row=56,column=c0+1,value=round(sum(e['valor'] for e in t['top'])/1e6,0)).number_format='#,##0'
        ws.cell(row=57,column=c0,value='Total declarado no ano').font=Font(bold=True)
        ws.cell(row=57,column=c0+1,value=round(t['total']/1e6,0)).number_format='#,##0'
        ws.cell(row=58,column=c0,value='% das 50 no total').font=Font(bold=True)
        ws.cell(row=58,column=c0+1,value=sum(e['valor'] for e in t['top'])/t['total']).number_format='0%'
    ws.column_dimensions[openpyxl.utils.get_column_letter(c0)].width=34
    ws.column_dimensions[openpyxl.utils.get_column_letter(c0+1)].width=9
ws.column_dimensions['A'].width=4
ws.freeze_panes='B5'
wb.save(X); print('ok')
for y in (2016,2017,2018):
    print(y,[ (e['nome'][:28],round(e['valor']/1e9,2)) for e in T[str(y)]['top'][:5]])
