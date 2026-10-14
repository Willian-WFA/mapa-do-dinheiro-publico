import json,os,collections,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
X='/media/hitsuzen/Novo volume/money/arrecadacao_despesas_2019_2026.xlsx'
wb=openpyxl.load_workbook(X)
if 'Gasto por área' in wb.sheetnames: del wb['Gasto por área']
ws=wb.create_sheet('Gasto por área',0)
hdr=PatternFill('solid',fgColor='1F3864'); grp=PatternFill('solid',fgColor='D9E1F2')
Y=list(range(2016,2027))
c=json.load(open('func_cache.json'))
a=collections.defaultdict(lambda:collections.defaultdict(float)); nm=collections.defaultdict(int)
for k,v in c.items():
    y=int(k[:4]); nm[y]+=1
    for f,(p,r) in v.items(): a[y][f]+=p+r
ben=json.load(open('ben.json')) if os.path.exists('ben.json') else {}
wbk=json.load(open('wb_enrl.json'))
pop={2016:206081432,2017:207660929,2018:208494900,2019:210147125,2020:211755692,2021:213317639,2024:212583750,2025:213421037}
pop[2022]=pop[2021]+(pop[2024]-pop[2021])/3; pop[2023]=pop[2021]+2*(pop[2024]-pop[2021])/3
planos={2016:47.7,2017:47.3,2018:47.3,2019:47.1,2020:47.6,2021:49.1,2022:50.5,2023:51.0,2024:52.1,2025:53.1}
ws['A1']='Gasto da União por área e pessoas atendidas, 2016–2026'; ws['A1'].font=Font(bold=True,size=14)
ws['A2']=('Gasto = despesa paga por função do orçamento da União (inclui pessoal e transferências a estados/municípios; soma valor pago no ano + restos a pagar pagos). '
          'Fonte: Portal da Transparência (despesas por ação). 2026 = janeiro a outubro. Não inclui gasto próprio de estados e municípios.')
ws['A2'].alignment=Alignment(wrap_text=True); ws['A2'].font=Font(italic=True,size=9); ws.row_dimensions[2].height=42
def head(r,t):
    for col,v in enumerate([t]+[str(y)+(' (jan–out)' if y==2026 else '') for y in Y],1):
        x=ws.cell(row=r,column=col,value=v); x.font=Font(bold=True,color='FFFFFF'); x.fill=hdr; x.alignment=Alignment(horizontal='center',wrap_text=True)
def row(r,label,vals,fmt='#,##0.0',bold=False,note=None):
    ws.cell(row=r,column=1,value=label).font=Font(bold=bold)
    for j,v in enumerate(vals):
        if v is not None:
            x=ws.cell(row=r,column=2+j,value=round(v,2)); x.number_format=fmt
    if note: ws.cell(row=r,column=len(Y)+2,value=note).font=Font(italic=True,size=9)
areas=[('Saúde','Saúde'),('Educação','Educação'),('Assistência social','Assistência social'),('Segurança pública','Segurança pública')]
others=[('Previdência social','Previdência social'),('Defesa nacional','Defesa nacional'),('Encargos especiais (dívida e transferências)','Encargos especiais')]
r=4; head(r,'GASTO DA UNIÃO (R$ bilhões)'); r+=1
for lab,f in areas: row(r,lab,[a[y][f]/1e9 for y in Y],bold=True); r+=1
r+=1; head(r,'Outras funções (para comparação)'); r+=1
for lab,f in others: row(r,lab,[a[y][f]/1e9 for y in Y]); r+=1
row(r,'Total pago pela União',[sum(a[y].values())/1e9 for y in Y],bold=True); r+=2
head(r,'GASTO POR HABITANTE (R$/ano)'); r+=1
for lab,f in areas:
    row(r,lab,[a[y][f]/pop[y] if y in pop else None for y in Y],'#,##0'); r+=1
row(r,'População (IBGE; 2022–2023 interpoladas)',[pop.get(y) for y in Y],'#,##0'); r+=2
head(r,'PESSOAS ATENDIDAS / BENEFICIADAS'); r+=1
def wbs(n,y): 
    v=wbk.get(n,{}).get(str(y)); return v/1e6 if v else None
row(r,'Educação – matrículas fund. + médio, todas as redes (milhões; Banco Mundial/UNESCO)',[ (wbs('Fundamental I (primário)',y)+wbs('Fund. II + Médio (secundário)',y)) if wbs('Fundamental I (primário)',y) and wbs('Fund. II + Médio (secundário)',y) else None for y in Y],'#,##0.0'); r+=1
row(r,'Educação – matrículas ensino superior (milhões; só até 2018)',[wbs('Superior',y) for y in Y],'#,##0.0'); r+=1
row(r,'Saúde – beneficiários de planos médico-hospitalares (milhões; ANS, aprox.)',[planos.get(y) for y in Y],'#,##0.0'); r+=1
row(r,'Saúde – população sem plano = dependente do SUS (milhões, estimativa)',[(pop[y]/1e6-planos[y]) if y in pop and y in planos else None for y in Y],'#,##0.0'); r+=1
def b(k,y): return ben.get(k,{}).get(str(y))
row(r,'Assistência – BPC (milhões de benefícios, dezembro)',[ (b('bpc',y)/1e6 if b('bpc',y) else None) for y in Y],'#,##0.00'); r+=1
row(r,'Assistência – Bolsa Família / Auxílio Brasil (milhões de famílias, dezembro)',[ (b('bf',y)/1e6 if b('bf',y) else None) for y in Y],'#,##0.00'); r+=1
row(r,'Segurança – sem beneficiário mensurável (serviço público universal)',[None]*len(Y)); r+=2
ws.cell(row=r,column=1,value='Observações').font=Font(bold=True); r+=1
for t in ['Saúde: o gasto federal inclui repasses a estados e municípios (SUS). O total nacional em saúde é maior (soma dos três níveis).',
          'Educação: não inclui o gasto de estados e municípios, que respondem pela maior parte da educação básica. Superior após 2018 e o Censo Escolar do INEP exigem o servidor do INEP (certificado não confiável nesta máquina).',
          'Segurança pública: a maior parte do gasto nacional é estadual; aqui só União (PF, PRF, Funpen, FNSP, Sistema Penitenciário Federal).',
          'Assistência social 2020 inclui o Auxílio Emergencial (R$ ~300 bi). 2021–2022 incluem Auxílio Brasil.',
          'Planos de saúde: valores aproximados de dezembro (ANS), a conferir; a ANS conta vínculos, não pessoas.']:
    ws.cell(row=r,column=1,value=t).alignment=Alignment(wrap_text=False); r+=1
ws.column_dimensions['A'].width=78
for j in range(len(Y)): ws.column_dimensions[openpyxl.utils.get_column_letter(2+j)].width=11
ws.column_dimensions[openpyxl.utils.get_column_letter(len(Y)+2)].width=40
ws.freeze_panes='B5'
wb.save(X); print('ok')
