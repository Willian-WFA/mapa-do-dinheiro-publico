import json,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
X='/media/hitsuzen/Novo volume/money/arrecadacao_despesas_2019_2026.xlsx'
C=json.load(open('compare.json')); G=C['groups']; Y=C['years']
bench=[g for g in G if g.get('bench')]; grp_=[g for g in G if not g.get('bench')]
wb=openpyxl.load_workbook(X)
if 'Quantas vezes mais' in wb.sheetnames: del wb['Quantas vezes mais']
ws=wb.create_sheet('Quantas vezes mais',3)
hdr=PatternFill('solid',fgColor='1F3864'); grp=PatternFill('solid',fgColor='D9E1F2')
ws['A1']='Quantas vezes mais cada grupo recebe por pessoa do que o gasto federal por pessoa atendida (educação, saúde, assistência, segurança)'; ws['A1'].font=Font(bold=True,size=13)
ws['A2']='Razão = valor por pessoa do grupo ÷ gasto federal por pessoa da referência (mesmo ano). Abaixo de cada bloco: quantas vezes o grupo de referência é maior que o grupo. Valores deflacionados não alteram a razão (mesmo ano). Isenções fiscais são renúncia, não pagamento.'; ws['A2'].font=Font(italic=True,size=9)
r=4
def pp(g,y):
    d=g['byYear'].get(str(y)); return (d[1]*1e9/d[0],d[0],d[2]) if d else None
for b in bench:
    ws.cell(row=r,column=1,value='REFERÊNCIA: '+b['label']+' ('+b['per']+')').font=Font(bold=True)
    for c in range(1,len(Y)+2): ws.cell(row=r,column=c).fill=grp
    r+=1
    ws.cell(row=r,column=1,value='Grupo (razão por pessoa)').font=Font(bold=True,color='FFFFFF'); ws.cell(row=r,column=1).fill=hdr
    for j,y in enumerate(Y):
        x=ws.cell(row=r,column=2+j,value=y); x.font=Font(bold=True,color='FFFFFF'); x.fill=hdr; x.alignment=Alignment(horizontal='center')
    r+=1
    # linha da referencia
    ws.cell(row=r,column=1,value='Referência: R$ por pessoa no ano')
    for j,y in enumerate(Y):
        p=pp(b,y)
        if p: ws.cell(row=r,column=2+j,value=round(p[0])).number_format='#,##0'
    r+=1
    ws.cell(row=r,column=1,value='Referência: nº de pessoas')
    for j,y in enumerate(Y):
        p=pp(b,y)
        if p: ws.cell(row=r,column=2+j,value=round(p[1])).number_format='#,##0'
    r+=1
    for g in grp_:
        ws.cell(row=r,column=1,value=g['label']+' – vezes mais por pessoa')
        for j,y in enumerate(Y):
            a=pp(g,y); p=pp(b,y)
            if a and p: ws.cell(row=r,column=2+j,value=round(a[0]/p[0],1)).number_format='#,##0.0'
        r+=1
        ws.cell(row=r,column=1,value='   quantas vezes o grupo de referência é maior que o grupo')
        for j,y in enumerate(Y):
            a=pp(g,y); p=pp(b,y)
            if a and p: ws.cell(row=r,column=2+j,value=round(p[1]/a[1],1)).number_format='#,##0.0'
        r+=1
    r+=1
ws.column_dimensions['A'].width=70
for j in range(len(Y)): ws.column_dimensions[openpyxl.utils.get_column_letter(2+j)].width=11
ws.freeze_panes='B4'
wb.save(X); print('ok')
