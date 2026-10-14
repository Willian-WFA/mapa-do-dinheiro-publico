import json,re,collections,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
X='/media/hitsuzen/Novo volume/money/arrecadacao_despesas_2019_2026.xlsx'
wb=openpyxl.load_workbook(X)
for n in ('Militares','Isenção e Agro'):
    if n in wb.sheetnames: del wb[n]
hdr=PatternFill('solid',fgColor='1F3864'); grp=PatternFill('solid',fgColor='D9E1F2')
def head(ws,r,vals):
    for c,v in enumerate(vals,1):
        x=ws.cell(row=r,column=c,value=v); x.font=Font(bold=True,color='FFFFFF'); x.fill=hdr; x.alignment=Alignment(horizontal='center',wrap_text=True,vertical='center')
# ---- Militares
cache=json.load(open('mil_cache.json'))
G={'Militares':'Ativa','Reserva_Reforma_Militares':'Reserva/Reforma','Pensionistas_DEFESA':'Pensionistas'}
d=collections.defaultdict(lambda:collections.defaultdict(dict))  # y -> g -> m -> (n,tot)
for k,v in cache.items():
    m=re.match(r'(\d{4})(\d{2})_(.+)\.zip',k); y,mm,o=int(m[1]),int(m[2]),m[3]
    d[y][G[o]][mm]=(v['n'],v['bruto_total'])
ws=wb.create_sheet('Militares')
ws['A1']='Militares do Executivo federal (Forças Armadas): quantidade e valor recebido por grupo e ano'; ws['A1'].font=Font(bold=True,size=13)
ws['A2']='Fonte: Portal da Transparência (Remuneração mensal). Valor = remuneração básica bruta + gratificação natalina + férias + eventuais, menos abate-teto; antes de IRRF/PSS. Quantidade = nº médio de pessoas por mês. Reserva e pensionistas só existem na base a partir de 2020.'; ws['A2'].font=Font(italic=True,size=9)
ws['A2'].alignment=Alignment(wrap_text=True); ws.row_dimensions[2].height=40
head(ws,4,['Ano','Meses na base']+sum([[f'{g} – pessoas (média/mês)',f'{g} – recebido no ano (R$ bi)',f'{g} – média por pessoa/mês (R$)'] for g in G.values()],[])+['Total recebido (R$ bi)'])
r=5
for y in sorted(d):
    meses=max(len(d[y][g]) for g in d[y]); row=[y,meses]; tot=0
    for g in G.values():
        ms=d[y].get(g,{})
        if ms:
            n=sum(v[0] for v in ms.values())/len(ms); t=sum(v[1] for v in ms.values())
            row+= [round(n),round(t/1e9,2),round(t/sum(v[0] for v in ms.values()),0)]; tot+=t
        else: row+=[None,None,None]
    row.append(round(tot/1e9,2))
    for c,v in enumerate(row,1):
        x=ws.cell(row=r,column=c,value=v)
        if c>2: x.number_format='#,##0.00' if c%3==0 or c==len(row) else '#,##0'
    if meses<12: ws.cell(row=r,column=len(row)+1,value=f'parcial ({meses} meses)')
    r+=1
for c in range(1,14): ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width=17
ws.row_dimensions[4].height=48
# ---- Isenção e Agro
T=json.load(open('ren_top.json')); A=json.load(open('agro.json')); S=json.load(open('rtn_sub.json'))
ws=wb.create_sheet('Isenção e Agro')
ws['A1']='Isenção fiscal: total, concentração e parcela do agro, por ano (R$ bilhões)'; ws['A1'].font=Font(bold=True,size=13)
ws['A2']='Fonte: Portal da Transparência – Renúncias Fiscais (CGU/Receita), autodeclarado. 2024 PARCIAL. Agro = CNAE 01–03 (agropecuária), 10–12 (agroindústria) e fertilizantes/defensivos/máquinas agrícolas. Isenções que não aparecem nesta base (IRPF rural, Funrural, parte do Simples) não estão incluídas.'
ws['A2'].font=Font(italic=True,size=9); ws['A2'].alignment=Alignment(wrap_text=True); ws.row_dimensions[2].height=40
yrs=list(range(2016,2025))
head(ws,4,['Item']+[str(y)+(' (parcial)' if y==2024 else '') for y in yrs])
def line(r,label,vals,fmt='#,##0.0',bold=False):
    ws.cell(row=r,column=1,value=label).font=Font(bold=bold)
    for j,v in enumerate(vals):
        c=ws.cell(row=r,column=2+j,value=None if v is None else round(v,2)); c.number_format=fmt
r=5
line(r,'Renúncia total declarada',[T[str(y)]['total']/1e9 for y in yrs],bold=True); r+=1
line(r,'Nº de empresas beneficiárias',[T[str(y)]['n'] for y in yrs],'#,##0'); r+=1
line(r,'15 maiores empresas – soma',[T[str(y)]['top15sum']/1e9 for y in yrs]); r+=1
line(r,'15 maiores – % do total',[T[str(y)]['top15sum']/T[str(y)]['total']*100 for y in yrs],'0.0'); r+=2
ws.cell(row=r,column=1,value='AGRO').font=Font(bold=True)
for c in range(1,11): ws.cell(row=r,column=c).fill=grp
r+=1
keys=['Agropecuária (CNAE 01–03)','Agroindústria (alimentos, bebidas, fumo)','Insumos e máquinas agrícolas']
for k in keys:
    line(r,k+' – valor',[A[str(y)]['agro'][k]['valor']/1e9 for y in yrs]); r+=1
    line(r,k+' – nº empresas',[A[str(y)]['agro'][k]['empresas'] for y in yrs],'#,##0'); r+=1
tot_agro=[sum(A[str(y)]['agro'][k]['valor'] for k in keys)/1e9 for y in yrs]
line(r,'Agro total (3 grupos)',tot_agro,bold=True); r+=1
line(r,'Agro – % da renúncia total',[tot_agro[i]/(T[str(y)]['total']/1e9)*100 for i,y in enumerate(yrs)],'0.0'); r+=2
ws.cell(row=r,column=1,value='SUBSÍDIOS E CRÉDITO RURAL – Orçamento da União (RTN, valor pago, R$ bi)').font=Font(bold=True)
for c in range(1,11): ws.cell(row=r,column=c).fill=grp
r+=1
for k,v in S.items():
    line(r,k,[v.get(str(y)) for y in yrs]); r+=1
r+=1
ws.cell(row=r,column=1,value='Tipos de renúncia mais relevantes em cada ano (R$ bi)').font=Font(bold=True)
for c in range(1,11): ws.cell(row=r,column=c).fill=grp
r+=1
for rank in range(5):
    ws.cell(row=r,column=1,value=f'#{rank+1}')
    for j,y in enumerate(yrs):
        t=A[str(y)]['tipo']
        if rank<len(t): ws.cell(row=r,column=2+j,value=f"{t[rank][0][:34]}: {t[rank][1]/1e9:.1f}")
    r+=1
r+=1
ws.cell(row=r,column=1,value='Maiores beneficiárias do AGRO (top 3 de cada grupo, R$ bi)').font=Font(bold=True)
for c in range(1,11): ws.cell(row=r,column=c).fill=grp
r+=1
for k in keys:
    for rank in range(3):
        ws.cell(row=r,column=1,value=f'{k[:22]} #{rank+1}')
        for j,y in enumerate(yrs):
            t=A[str(y)]['agro'][k]['top']
            if rank<len(t): ws.cell(row=r,column=2+j,value=f"{t[rank][0].title()[:26]}: {t[rank][1]/1e9:.2f}")
        r+=1
ws.column_dimensions['A'].width=52
for j in range(len(yrs)): ws.column_dimensions[openpyxl.utils.get_column_letter(2+j)].width=34 if False else 22
wb.save(X)
print('ok')
for y in sorted(d):
    print(y,{g:(round(sum(v[0] for v in d[y][g].values())/len(d[y][g])),round(sum(v[1] for v in d[y][g].values())/1e9,1),len(d[y][g])) for g in d[y]})
