import json,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
X='/media/hitsuzen/Novo volume/money/arrecadacao_despesas_2019_2026.xlsx'
H=json.load(open('hist_rows.json'))
wb=openpyxl.load_workbook(X)
if 'Pessoas por benefício' in wb.sheetnames: del wb['Pessoas por benefício']
ws=wb.create_sheet('Pessoas por benefício',2)
hdr=PatternFill('solid',fgColor='1F3864')
ws['A1']='Quantas pessoas, famílias ou empresas em cada benefício / grupo'; ws['A1'].font=Font(bold=True,size=14)
ws['A2']='Último ano disponível de cada série e variação desde o primeiro ano. Milhões, mil ou número, conforme o rótulo. Séries completas na aba "Série histórica".'; ws['A2'].font=Font(italic=True,size=9)
for c,v in enumerate(['Grupo / benefício','Primeiro ano','Valor','Último ano','Valor','Variação'],1):
    x=ws.cell(row=4,column=c,value=v); x.fill=hdr; x.font=Font(bold=True,color='FFFFFF')
r=5
for row in H['rows']:
    if not row['sec'].startswith('2)'): continue
    d={int(y):v for y,v in row['data'].items() if v is not None}
    if not d: continue
    ys=sorted(d); a,b=ys[0],ys[-1]
    ws.cell(row=r,column=1,value=row['label'].strip())
    ws.cell(row=r,column=2,value=a); ws.cell(row=r,column=3,value=round(d[a],2)).number_format='#,##0.00'
    ws.cell(row=r,column=4,value=b); ws.cell(row=r,column=5,value=round(d[b],2)).number_format='#,##0.00'
    if d[a]: ws.cell(row=r,column=6,value=d[b]/d[a]-1).number_format='+0%;-0%'
    r+=1
ws.column_dimensions['A'].width=78
for c in 'BCDEF': ws.column_dimensions[c].width=13
wb.save(X); print('ok',r-5)
