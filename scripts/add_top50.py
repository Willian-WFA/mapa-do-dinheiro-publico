import json,openpyxl,collections
from openpyxl.styles import Font,PatternFill,Alignment
X='/media/hitsuzen/Novo volume/money/arrecadacao_despesas_2019_2026.xlsx'
T=json.load(open('top50.json')); Y=list(range(2016,2025))
wb=openpyxl.load_workbook(X)
for n in ('Top50 isenção','Top50 por empresa'):
    if n in wb.sheetnames: del wb[n]
hdr=PatternFill('solid',fgColor='1F3864')
def H(ws,r,vals):
    for c,v in enumerate(vals,1):
        x=ws.cell(row=r,column=c,value=v); x.fill=hdr; x.font=Font(bold=True,color='FFFFFF'); x.alignment=Alignment(horizontal='center',wrap_text=True)
# ---- lista longa
ws=wb.create_sheet('Top50 isenção',wb.sheetnames.index('Top15 isenção')+1 if 'Top15 isenção' in wb.sheetnames else None)
ws['A1']='50 empresas com maior isenção/renúncia fiscal por ano (matriz + filiais somadas pela raiz do CNPJ)'; ws['A1'].font=Font(bold=True,size=14)
ws['A2']='Fonte: Portal da Transparência – Renúncias Fiscais (CGU/Receita Federal), valores autodeclarados. 2024 é PARCIAL (cerca de 10 mil beneficiários). Valores nominais, R$ milhões. Setor = CNAE da maior linha da empresa; "Sem informação" é comum.'
ws['A2'].font=Font(italic=True,size=9); ws['A2'].alignment=Alignment(wrap_text=True); ws.row_dimensions[2].height=32
H(ws,4,['Ano','#','Empresa','CNPJ (raiz)','Setor (CNAE)','UF','Renúncia (R$ mi)','% do total do ano','Acumulado top 50 (%)','Principal tipo de benefício','% desse tipo'])
r=5
for y in Y:
    t=T[str(y)]; acc=0
    for i,e in enumerate(t['top'],1):
        acc+=e['valor']
        vals=[y,i,e['nome'].title()[:50],e['cnpj'],e['cnae'][:50],e['uf'],e['valor']/1e6,e['valor']/t['total'],acc/t['total'],e['tipo'][:60],e['tipo_pct']]
        for c,v in enumerate(vals,1):
            x=ws.cell(row=r,column=c,value=v)
            if c==7: x.number_format='#,##0.0'
            if c in (8,9,11): x.number_format='0.0%'
        r+=1
for c,w in zip('ABCDEFGHIJK',[6,5,46,12,40,5,15,12,14,48,10]): ws.column_dimensions[c].width=w
ws.freeze_panes='D5'; ws.auto_filter.ref=f'A4:K{r-1}'
# ---- matriz por empresa
agg=collections.defaultdict(lambda:{'nome':'','cnae':'','uf':'','v':{}, 'rank':{}})
for y in Y:
    for i,e in enumerate(T[str(y)]['top'],1):
        a=agg[e['cnpj']]; a['nome']=e['nome']; a['cnae']=e['cnae']; a['uf']=e['uf']; a['v'][y]=e['valor']; a['rank'][y]=i
ws=wb.create_sheet('Top50 por empresa',wb.sheetnames.index('Top50 isenção')+1)
ws['A1']='Empresas que aparecem entre as 50 maiores em algum ano – valor por ano (R$ milhões)'; ws['A1'].font=Font(bold=True,size=14)
ws['A2']='Ordenado pelo total acumulado nos anos em que apareceu no top 50. Vazio = fora do top 50 naquele ano. 2024 parcial.'; ws['A2'].font=Font(italic=True,size=9)
H(ws,4,['Empresa','CNPJ (raiz)','Setor (CNAE)','UF']+[str(y)+(' (parc.)' if y==2024 else '') for y in Y]+['Total no top 50','Anos no top 50','Melhor posição'])
rows=sorted(agg.items(),key=lambda kv:-sum(kv[1]['v'].values()))
r=5
for cnpj,a in rows:
    vals=[a['nome'].title()[:50],cnpj,a['cnae'][:45],a['uf']]+[ (a['v'][y]/1e6 if y in a['v'] else None) for y in Y]+[sum(a['v'].values())/1e6,len(a['v']),min(a['rank'].values())]
    for c,v in enumerate(vals,1):
        x=ws.cell(row=r,column=c,value=v)
        if 5<=c<=14: x.number_format='#,##0.0'
    r+=1
ws.column_dimensions['A'].width=46; ws.column_dimensions['B'].width=12; ws.column_dimensions['C'].width=36; ws.column_dimensions['D'].width=5
for c in range(5,15): ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width=11
ws.freeze_panes='E5'; ws.auto_filter.ref=f'A4:O{r-1}'
wb.save(X); print('ok',len(rows),'empresas distintas')
for cnpj,a in rows[:12]: print(a['nome'][:40].ljust(40),round(sum(a['v'].values())/1e9,1),len(a['v']),a['cnae'][:30])
